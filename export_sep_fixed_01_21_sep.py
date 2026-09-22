# -*- coding: utf-8 -*-
"""
export_sep_fixed_01_21_sep.py
Generates full Tracking Status Logs Report for All Bills:
  - Date Range: 01/09/2026 to 21/09/2026 (00:00:00 to 23:59:59)
  - Sheets:
      1. "All Tracking Logs" (All bills Day 1 to Day 21 + Aug carryover)
      2. "01 Sep" to "21 Sep" (Bills created on each specific day)
      3. "Aug Carryover" (Bills created before Sep but active/delivered in Sep)
  - 229 columns format identical to fixed specifications
"""

import os
import sys
import json
import sqlite3
import shutil
import requests
import pandas as pd
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(HERE, "config.json")
CACHE_DB_PATH = os.path.join(HERE, "cache", "tracking_cache_wide_sep_01_21.db")
TERMINAL_STATUSES = {"410", "520", "201", "600", "540"}

FULL_STATUS_CODES_TEMPLATE = [
    "110", "120", "130", "140", "150",
    "200", "201", "202", "205", "210", "220", "230", "240", "250",
    "300", "302", "304", "306", "308", "309", "310", "311", "312", "315", "320", "330",
    "400", "401", "402", "403", "404", "405", "406", "408", "415", "417", "420", "421", "422", "425", "428",
    "430", "431", "432", "440", "450", "460", "470", "471", "472", "475", "480", "485", "490", "495",
    "500", "501", "502", "505", "510", "512", "515", "520", "525", "530", "540", "550", "560", "570", "580", "590",
    "600", "410", "99"
]

def init_cache_db():
    os.makedirs(os.path.dirname(CACHE_DB_PATH), exist_ok=True)
    conn = sqlite3.connect(CACHE_DB_PATH)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS tracking_cache (
            order_id TEXT PRIMARY KEY,
            data_json TEXT
        )
    """)
    conn.commit()
    conn.close()

def load_cached_orders():
    init_cache_db()
    conn = sqlite3.connect(CACHE_DB_PATH)
    cached = {}
    for r in conn.execute("SELECT order_id, data_json FROM tracking_cache").fetchall():
        try:
            cached[r[0]] = json.loads(r[1])
        except Exception:
            pass
    conn.close()
    return cached

def save_cache_batch(items):
    if not items:
        return
    conn = sqlite3.connect(CACHE_DB_PATH)
    conn.executemany("INSERT OR REPLACE INTO tracking_cache (order_id, data_json) VALUES (?, ?)", items)
    conn.commit()
    conn.close()

def write_sheet_data(ws, sheet_title, rows, col_headers, styles):
    f_title, f_header, f_data, f_log, f_vas, fill_title, fill_hdr1, fill_hdr2, fill_alt, border, align_c, align_l = styles

    # Row 1: Title
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(col_headers))
    t_cell = ws.cell(1, 1, f"{sheet_title} ({len(rows)} Bills)")
    t_cell.font = f_title
    t_cell.fill = fill_title
    t_cell.alignment = align_c
    ws.row_dimensions[1].height = 30

    # Row 3: Headers
    ws.row_dimensions[3].height = 25
    for c_idx, h_text in enumerate(col_headers, 1):
        cell = ws.cell(3, c_idx, h_text)
        cell.font = f_header
        cell.fill = fill_hdr1 if c_idx in (1, 2, 3) or ((c_idx - 4) // 3) % 2 == 0 else fill_hdr2
        cell.alignment = align_c
        cell.border = border

    # Rows 4+: Data
    r_idx = 4
    for r_data in rows:
        row_fill = fill_alt if r_idx % 2 == 0 else None
        ws.row_dimensions[r_idx].height = 19

        c_bill = ws.cell(r_idx, 1, str(r_data.get("BILL ID", "")))
        c_bill.font = f_data; c_bill.border = border
        c_bill.alignment = align_l

        c_svc = ws.cell(r_idx, 2, str(r_data.get("SERVICE TYPE", "")))
        c_svc.font = f_data; c_svc.border = border
        if row_fill: c_svc.fill = row_fill
        c_svc.alignment = align_c

        c_vas = ws.cell(r_idx, 3, str(r_data.get("VAS", "")))
        c_vas.font = f_vas if r_data.get("VAS") else f_data
        c_vas.border = border
        if row_fill: c_vas.fill = row_fill
        c_vas.alignment = align_c

        trips_map = r_data.get("_trips_map", {})
        col_pos = 4
        for sc in FULL_STATUS_CODES_TEMPLATE:
            unit_val, time_val = trips_map.get(sc, ("", ""))
            log_val = sc if unit_val or time_val else ""

            c_log = ws.cell(r_idx, col_pos, log_val)
            c_log.font = f_log; c_log.border = border
            if row_fill: c_log.fill = row_fill
            c_log.alignment = align_c
            col_pos += 1

            c_unit = ws.cell(r_idx, col_pos, unit_val)
            c_unit.font = f_data; c_unit.border = border
            if row_fill: c_unit.fill = row_fill
            c_unit.alignment = align_c
            col_pos += 1

            c_time = ws.cell(r_idx, col_pos, time_val)
            c_time.font = f_data; c_time.border = border
            if row_fill: c_time.fill = row_fill
            c_time.alignment = align_c
            col_pos += 1

        for val in [r_data.get("LATEST_STATUS",""), r_data.get("LATEST_UNIT",""), r_data.get("LATEST_TIME",""), r_data.get("LATEST_USER","")]:
            c_last = ws.cell(r_idx, col_pos, val)
            c_last.font = f_data; c_last.border = border
            if row_fill: c_last.fill = row_fill
            c_last.alignment = align_c
            col_pos += 1

        r_idx += 1

    ws.freeze_panes = "D4"

    # Set column widths efficiently (fixed standard widths for predictable performance)
    for c_idx in range(1, len(col_headers) + 1):
        col_letter = get_column_letter(c_idx)
        if c_idx == 1:
            ws.column_dimensions[col_letter].width = 15
        elif c_idx == 2:
            ws.column_dimensions[col_letter].width = 15
        elif c_idx == 3:
            ws.column_dimensions[col_letter].width = 10
        elif c_idx >= len(col_headers) - 3:
            ws.column_dimensions[col_letter].width = 18 if c_idx == len(col_headers) - 1 else 14
        else:
            sub_col = (c_idx - 4) % 3
            if sub_col == 0:
                ws.column_dimensions[col_letter].width = 11
            elif sub_col == 1:
                ws.column_dimensions[col_letter].width = 13
            else:
                ws.column_dimensions[col_letter].width = 18

def main():
    sys.stdout.reconfigure(encoding="utf-8")
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        cfg = json.load(f)

    # Check detail file
    detail_file = os.path.join(HERE, "cache", "latest_detail.xlsx")
    if not os.path.exists(detail_file):
        print(f"[INFO] Downloading fresh detail file from API...")
        import downloader
        downloader.download_detail(cfg["api"], detail_file, from_date="20260901", to_date="20260921", force_refresh=True)

    print(f"[INFO] Loading detail data: {detail_file}...")
    try:
        df = pd.read_excel(detail_file, engine="calamine")
    except Exception:
        df = pd.read_excel(detail_file)
    df.columns = [str(c).strip().upper() for c in df.columns]
    print(f"[INFO] Total rows loaded: {len(df)}")

    # Filter: CURRENT TIME or CREATED DATE must be within Sep 1 - Sep 21 2026
    time_col = next((c for c in df.columns if "CURRENT TIME" in c), None)
    created_col = next((c for c in df.columns if "CREATED DATE" in c), None)
    if not time_col:
        print("[ERROR] CURRENT TIME column not found!")
        return

    df["_parsed_cur_time"] = pd.to_datetime(df[time_col], dayfirst=True, errors='coerce')
    df["_parsed_created"]  = pd.to_datetime(df[created_col], dayfirst=True, errors='coerce') if created_col else None

    sep_start = pd.Timestamp("2026-09-01 00:00:00")
    sep_end   = pd.Timestamp("2026-09-21 23:59:59")

    mask = (df["_parsed_cur_time"] >= sep_start) & (df["_parsed_cur_time"] <= sep_end)
    if created_col:
        mask = mask | ((df["_parsed_created"] >= sep_start) & (df["_parsed_created"] <= sep_end))

    df_filtered = df[mask].copy()
    print(f"[INFO] Bills in Sep 1-21 (00:00:00 to 23:59:59): {len(df_filtered)}")

    # Map order id to its creation day and current time day
    bill_day_group = {}
    for _, r in df_filtered.iterrows():
        oid = str(r.get("ORDER ID", "")).strip()
        c_dt = r.get("_parsed_created")
        t_dt = r.get("_parsed_cur_time")
        if pd.notna(c_dt) and c_dt.month == 9 and 1 <= c_dt.day <= 21:
            bill_day_group[oid] = f"{c_dt.day:02d} Sep"
        elif pd.notna(c_dt) and c_dt.month == 8:
            bill_day_group[oid] = "Aug Carryover"
        elif pd.notna(t_dt) and t_dt.month == 9 and 1 <= t_dt.day <= 21:
            bill_day_group[oid] = f"{t_dt.day:02d} Sep"
        else:
            bill_day_group[oid] = "Aug Carryover"

    # Build metadata maps
    service_map = {}
    svc_col = next((c for c in df_filtered.columns if c in ["SERVICE", "SERVICE TYPE", "SERVICE_TYPE", "SERVICE NAME", "SERVICETYPE"]), None)
    if svc_col:
        for _, r in df_filtered.iterrows():
            oid = str(r.get("ORDER ID", "")).strip()
            val = str(r.get(svc_col, "") or "").strip()
            if oid and val and val.lower() != "nan":
                service_map[oid] = val

    action_user_map = {}
    au_col = next((c for c in df_filtered.columns if c in ["ACTION USER", "ACTION_USER", "LAST ACTION USER", "LAST USER", "USER"]), None)
    if au_col:
        for _, r in df_filtered.iterrows():
            oid = str(r.get("ORDER ID", "")).strip()
            val = str(r.get(au_col, "") or "").strip()
            if oid and val and val.lower() != "nan":
                action_user_map[oid] = val

    vas_fee_map = {}
    vf_col = next((c for c in df_filtered.columns if "VAS FEE" in c), None)
    if vf_col:
        for _, r in df_filtered.iterrows():
            oid = str(r.get("ORDER ID", "")).strip()
            val = r.get(vf_col, 0)
            try:
                if float(val) > 0:
                    vas_fee_map[oid] = float(val)
            except Exception:
                pass

    current_status_map = {}
    st_col = next((c for c in df_filtered.columns if "CURRENT STATUS" in c), None)
    if st_col:
        for _, r in df_filtered.iterrows():
            oid = str(r.get("ORDER ID", "")).strip()
            val = str(r.get(st_col, "") or "").strip()
            if oid and val and val.lower() != "nan":
                current_status_map[oid] = val

    current_po_map = {}
    po_col = next((c for c in df_filtered.columns if "CURRENT POST OFFICE" in c), None)
    if po_col:
        for _, r in df_filtered.iterrows():
            oid = str(r.get("ORDER ID", "")).strip()
            val = str(r.get(po_col, "") or "").strip()
            if oid and val and val.lower() != "nan":
                current_po_map[oid] = val

    current_time_map = {}
    if time_col:
        for _, r in df_filtered.iterrows():
            oid = str(r.get("ORDER ID", "")).strip()
            val = r.get(time_col)
            if oid and pd.notna(val):
                try:
                    current_time_map[oid] = val.strftime("%d/%m/%Y %H:%M:%S")
                except Exception:
                    current_time_map[oid] = str(val)

    # Cache management
    cached_results = load_cached_orders()
    unique_oids = df_filtered["ORDER ID"].dropna().astype(str).unique()
    print(f"\n[INFO] Unique bills to process: {len(unique_oids)}")
    print(f"[INFO] Already cached: {len(cached_results)}")

    to_fetch = [oid for oid in unique_oids if oid not in cached_results]
    print(f"[INFO] Need to fetch from API: {len(to_fetch)}")

    headers_api = {
        "Authorization": "Bearer " + cfg["api"]["bearer_token"],
        "Accept": "application/json, text/plain, */*",
        "User-Agent": "Mozilla/5.0"
    }
    session = requests.Session()
    adapter = HTTPAdapter(pool_connections=100, pool_maxsize=100, max_retries=Retry(total=2, backoff_factor=0.1))
    session.mount("https://", adapter)
    session.headers.update(headers_api)

    def fetch_bill(oid_str):
        if not oid_str or oid_str.lower() == "nan":
            return None

        svc_type  = service_map.get(oid_str, "")
        raw_st    = current_status_map.get(oid_str, "")
        def_st    = raw_st.split("-")[0].strip() if raw_st else ""
        def_unit  = current_po_map.get(oid_str, "")
        def_time  = current_time_map.get(oid_str, "")
        def_user  = action_user_map.get(oid_str, "")

        row_dict = {
            "BILL ID": oid_str, "SERVICE TYPE": svc_type, "VAS": "",
            "_trips_map": {},
            "LATEST_STATUS": def_st, "LATEST_UNIT": def_unit,
            "LATEST_TIME": def_time, "LATEST_USER": def_user,
            "_sheet_group": bill_day_group.get(oid_str, "All Tracking Logs")
        }

        trips = []
        data_tr = None
        try:
            r_tr = session.get(
                "https://gw-express.metfone.com.kh/tms-tracking/api/v1/order-tracking",
                params={"order_id": oid_str}, timeout=12
            )
            if r_tr.status_code == 200:
                data_tr = r_tr.json()
                trips = data_tr.get("trackingTrips", [])
        except Exception:
            pass

        vas_val = ""
        try:
            r_sc = session.get(
                "https://gw-express.metfone.com.kh/tms-receiving/api/v1/orders/search",
                params={"order_code": oid_str}, timeout=8
            )
            if r_sc.status_code == 200:
                vas_val = str(r_sc.json().get("added_service_code") or "").strip()
        except Exception:
            pass

        if not vas_val and oid_str in vas_fee_map:
            vas_val = "VTT"
        row_dict["VAS"] = vas_val

        if not trips:
            return row_dict

        trips_sorted = list(reversed(trips))
        if not svc_type and isinstance(data_tr, dict):
            svc_type = str(data_tr.get("serviceType") or data_tr.get("serviceName") or data_tr.get("service") or "").strip()
            row_dict["SERVICE TYPE"] = svc_type

        latest_st = latest_unit = latest_time = latest_user = ""
        for t in trips_sorted:
            st = str(t.get("status", "") or "").lstrip("S").strip()
            po = t.get("postOffice") or {}
            unit = t.get("postcode") or (po.get("code") if isinstance(po, dict) else "") or ""
            if not unit and "handoverInfo" in t:
                unit = t.get("handoverInfo", {}).get("departmentCode", "")

            dt_raw = t.get("updatedAt", "")
            dt_str = ""
            dt_obj = None
            if dt_raw:
                try:
                    dt_obj = pd.to_datetime(dt_raw)
                    dt_str = dt_obj.strftime("%d/%m/%Y %H:%M:%S")
                except Exception:
                    dt_str = str(dt_raw)

            if dt_obj is not None:
                dt_naive = dt_obj.replace(tzinfo=None) if dt_obj.tzinfo else dt_obj
                if not (sep_start <= dt_naive <= sep_end):
                    continue

            upd_obj = t.get("updatedBy")
            upd_user = ""
            if isinstance(upd_obj, dict):
                upd_user = str(upd_obj.get("name") or "").strip()
            elif isinstance(upd_obj, str):
                upd_user = upd_obj.strip()

            if st:
                row_dict["_trips_map"][st] = (unit, dt_str)
            latest_st = st
            latest_unit = unit
            latest_time = dt_str
            if upd_user:
                latest_user = upd_user

        if not latest_user:
            latest_user = action_user_map.get(oid_str, "")

        row_dict["LATEST_STATUS"] = latest_st or def_st
        row_dict["LATEST_UNIT"]   = latest_unit or def_unit
        row_dict["LATEST_TIME"]   = latest_time or def_time
        row_dict["LATEST_USER"]   = latest_user or def_user
        return row_dict

    results_dict = dict(cached_results)
    if to_fetch:
        print(f"\n[INFO] Fetching {len(to_fetch)} bills with 80 threads...", flush=True)
        batch_to_save = []
        count = 0
        total = len(to_fetch)
        with ThreadPoolExecutor(max_workers=80) as executor:
            futures = {executor.submit(fetch_bill, oid): oid for oid in to_fetch}
            for future in as_completed(futures):
                oid = futures[future]
                try:
                    res = future.result()
                    if res:
                        results_dict[oid] = res
                        batch_to_save.append((oid, json.dumps(res, ensure_ascii=False)))
                except Exception:
                    pass
                count += 1
                if count % 1000 == 0 or count == total:
                    save_cache_batch(batch_to_save)
                    batch_to_save = []
                    print(f"[PROGRESS] {count}/{total} fetched ({len(results_dict)} total)...", flush=True)
        if batch_to_save:
            save_cache_batch(batch_to_save)

    all_results = []
    for oid in unique_oids:
        if oid in results_dict:
            res = results_dict[oid]
            res["_sheet_group"] = bill_day_group.get(oid, "Aug Carryover")
            all_results.append(res)

    print(f"\n[INFO] Total bills for Excel: {len(all_results)}", flush=True)

    # Prepare sheets dictionary
    sheets_data = {"All Tracking Logs": all_results}
    
    # Pre-populate daily sheets for Day 1 to Day 21
    for d in range(1, 22):
        d_name = f"{d:02d} Sep"
        sheets_data[d_name] = []
    sheets_data["Aug Carryover"] = []

    for r in all_results:
        grp = r.get("_sheet_group", "Aug Carryover")
        if grp in sheets_data:
            sheets_data[grp].append(r)
        else:
            sheets_data["Aug Carryover"].append(r)

    print("\n[INFO] Breakdown across sheets:", flush=True)
    for s_name, s_rows in sheets_data.items():
        print(f"  - Sheet '{s_name}': {len(s_rows)} bills", flush=True)

    # Build Excel Workbook
    wb = Workbook()
    ws_first = wb.active

    font_family = "Calibri"
    f_title  = Font(name=font_family, size=13, bold=True, color="FFFFFF")
    f_header = Font(name=font_family, size=10, bold=True, color="FFFFFF")
    f_data   = Font(name=font_family, size=9)
    f_log    = Font(name=font_family, size=9, bold=True, color="1E3A8A")
    f_vas    = Font(name=font_family, size=9, bold=True, color="047857")

    fill_title = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_hdr1  = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_hdr2  = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    fill_alt   = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin       = Side(border_style="thin", color="CBD5E1")
    border     = Border(left=thin, right=thin, top=thin, bottom=thin)
    align_c    = Alignment(horizontal="center", vertical="center")
    align_l    = Alignment(horizontal="left", vertical="center")

    styles = (f_title, f_header, f_data, f_log, f_vas, fill_title, fill_hdr1, fill_hdr2, fill_alt, border, align_c, align_l)

    col_headers = ["BILL ID", "SERVICE TYPE", "VAS"]
    for sc in FULL_STATUS_CODES_TEMPLATE:
        col_headers.extend([f"LOG {sc}", f"UNIT {sc}", f"TIME {sc}"])
    col_headers.extend(["LATEST STATUS", "LATEST UNIT", "LATEST TIME", "LAST USER"])

    first_done = False
    for s_name, s_rows in sheets_data.items():
        if not s_rows and s_name != "All Tracking Logs":
            continue
        
        s_rows.sort(key=lambda x: str(x.get("BILL ID", "")))
        if not first_done:
            ws = ws_first
            ws.title = s_name
            first_done = True
        else:
            ws = wb.create_sheet(title=s_name)

        if s_name == "All Tracking Logs":
            title_text = "ALL BILL TRACKING STATUS LOGS REPORT - 01 TO 21 SEPTEMBER 2026"
        elif s_name == "Aug Carryover":
            title_text = "BILL TRACKING STATUS LOGS - AUGUST CARRYOVER (ACTIVE IN SEP 2026)"
        else:
            title_text = f"BILL TRACKING STATUS LOGS - {s_name.upper()} 2026"

        print(f"[INFO] Writing sheet '{s_name}' ({len(s_rows)} rows)...", flush=True)
        write_sheet_data(ws, title_text, s_rows, col_headers, styles)

    primary_new = r"C:\Users\DELL\Desktop\Bill_Tracking_Status_Logs_01Sep_21Sep_20260921_FIXED.xlsx"
    local_new   = os.path.join(HERE, "Bill_Tracking_Status_Logs_01Sep_21Sep_20260921_FIXED.xlsx")

    print(f"\n[INFO] Saving workbook to: {primary_new}...", flush=True)
    wb.save(primary_new)
    print(f"[SUCCESS] Saved primary desktop file: {primary_new}", flush=True)

    try:
        shutil.copy2(primary_new, local_new)
        print(f"[SUCCESS] Synced report to: {local_new}", flush=True)
    except Exception as e:
        print(f"[WARNING] Could not sync to {local_new}: {e}", flush=True)

    print(f"\n[COMPLETE] Report ready - {len(all_results)} bills across {len(sheets_data)} sheets (Day 1-21 + Master)!", flush=True)

if __name__ == "__main__":
    main()
